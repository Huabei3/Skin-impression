"""
单独的测试/推理脚本 —— 加载 best_model.pth 在 test set 上跑评估，输出 xlsx。

用法（实例2 v3 SA）:
  python predict_p/test_only.py --model-variant v3 --rgb-backbone simple_cnn \
      --global-backbone simple_cnn --race SA \
      --data-root /root/autodl-tmp \
      --gt-excel /root/autodl-tmp/gt/toMax_gt.xlsx \
      --output-root /root/autodl-tmp/deepskin/predict_p/output

输出 xlsx 包含:
  - Params sheet: 推理参数
  - Overall sheet: 整体指标 (MAE, Pearson, etc.)
  - PerGroup sheet: 每个 model+ior+scene 的 pearson_r
"""
from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Dict, List

import numpy as np
import torch
from torch.cuda.amp import autocast

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from predict_p.config import Config
from predict_p.trainer import Trainer
from predict_p.train import _apply_model_variant_overrides, _apply_race_overrides
from predict_p.metrics import evaluate_scores

logger = logging.getLogger(__name__)


def _extract_model_ior_scene(original_name: str) -> tuple:
    """
    从 original_name 提取 (model, ior, scene)。

    original_name 示例:
      "f07ih3k_01"  → model="f07", ior="i", scene="h3k"
      "f08rrs02_15" → model="f08", ior="r", scene="rs02"
      "m08rh4k_03"  → model="m08", ior="r", scene="h4k"

    规则: 前3字符=model, 第4字符=ior, 剩余到 _ 之前=scene
    """
    s = str(original_name)
    # 去掉末尾 _数字 后缀得到前缀
    parts = s.rsplit("_", 1)
    prefix = parts[0] if len(parts) == 2 and parts[1].isdigit() else s

    if len(prefix) >= 4:
        model = prefix[:3]
        ior = prefix[3]
        scene = prefix[4:] if len(prefix) > 4 else ""
    else:
        model = prefix
        ior = ""
        scene = ""
    return model, ior, scene


def _maybe_normalize_target_score(target_score: torch.Tensor) -> torch.Tensor:
    if target_score.max() > 1.5:
        return target_score / 10.0
    return target_score


def run_test_with_metadata(trainer: Trainer, ckpt_path: str = None) -> Dict:
    """
    跑推理并收集每个样本的 original_name、预测值、GT 值。
    返回 dict，包含 overall metrics 和 per-sample 列表。
    """
    if ckpt_path is not None:
        ckpt_path = Path(ckpt_path)
    else:
        ckpt_path = Path(trainer.config["CHECKPOINT_DIR"]) / "best_model.pth"

    if ckpt_path.exists():
        ckpt = torch.load(ckpt_path, map_location=trainer.device)
        trainer.model.load_state_dict(ckpt["model_state_dict"], strict=True)
        logger.info(f"Loaded checkpoint from {ckpt_path}")
    else:
        logger.warning(f"Checkpoint not found: {ckpt_path}, using current model weights")

    trainer.model.eval()

    all_pred = []
    all_tgt = []
    all_meta = []  # 每个样本的 original_name, scene_person

    scaler_ctx = autocast() if trainer.scaler else torch.no_grad()

    with torch.no_grad():
        for batch in trainer.test_loader:
            face_rgb = batch["face_rgb"].to(trainer.device, non_blocking=True)
            global_rgb = batch.get("global_rgb", face_rgb).to(trainer.device, non_blocking=True)
            target_score = batch["preference_score"].to(trainer.device, non_blocking=True)
            face_uv = None
            if str(trainer.config.get("MODEL_VARIANT", "v1")).lower().strip() != "v3":
                face_uv = batch["face_uv"].to(trainer.device, non_blocking=True)

            with scaler_ctx:
                pred_logits = trainer.model(face_rgb, face_uv, global_rgb)

            all_pred.append(torch.sigmoid(pred_logits).cpu())
            all_tgt.append(_maybe_normalize_target_score(target_score).cpu())

    # 现在从 test_loader.dataset 收集 metadata（顺序一致）
    test_items = trainer.test_loader.dataset.data_items
    for item in test_items:
        all_meta.append({
            "original_name": str(item.get("original_name", "")),
            "scene_person": str(item.get("scene_person", "")),
        })

    ps = torch.cat(all_pred, dim=0).numpy().flatten() if all_pred else np.array([])
    ts = torch.cat(all_tgt, dim=0).numpy().flatten() if all_tgt else np.array([])

    # 确保长度一致
    n = min(len(ps), len(ts), len(all_meta))
    ps = ps[:n]
    ts = ts[:n]
    all_meta = all_meta[:n]

    overall_metrics = evaluate_scores(ps, ts)

    return {
        "overall": overall_metrics,
        "predictions": ps.tolist(),
        "targets": ts.tolist(),
        "metadata": all_meta,
    }


def compute_per_group_pearson(ps: np.ndarray, ts: np.ndarray, metadata: list) -> List[Dict]:
    """按 model+ior+scene 分组计算 pearson_r。"""
    groups: Dict[str, Dict[str, list]] = {}

    for i, meta in enumerate(metadata):
        model, ior, scene = _extract_model_ior_scene(meta["original_name"])
        key = f"{model}_{ior}_{scene}"

        if key not in groups:
            groups[key] = {"model": model, "ior": ior, "scene": scene, "pred": [], "tgt": []}
        groups[key]["pred"].append(float(ps[i]))
        groups[key]["tgt"].append(float(ts[i]))

    results = []
    for key in sorted(groups.keys()):
        g = groups[key]
        pred_arr = np.array(g["pred"])
        tgt_arr = np.array(g["tgt"])
        n = len(pred_arr)

        pearson_r = float("nan")
        if n > 1:
            try:
                pearson_r = float(np.corrcoef(pred_arr, tgt_arr)[0, 1])
            except Exception:
                pass

        mae = float(np.mean(np.abs(pred_arr - tgt_arr))) if n > 0 else float("nan")

        results.append({
            "model": g["model"],
            "ior": g["ior"],
            "scene": g["scene"],
            "n_samples": n,
            "pearson_r": pearson_r,
            "mae": mae,
        })

    return results


def export_to_xlsx(
    output_path: str,
    args: argparse.Namespace,
    overall: Dict[str, float],
    per_group: List[Dict],
    metadata: list,
    predictions: list,
    targets: list,
) -> None:
    """导出 xlsx，包含 Params / Overall / PerGroup / RawData 四个 sheet。"""
    import openpyxl as xl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = xl.Workbook()

    # ----- styles -----
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def write_header(ws, headers, row=1):
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    # ==================== Sheet 1: Params ====================
    ws1 = wb.active
    ws1.title = "Params"
    ws1.cell(row=1, column=1, value="Parameter").font = header_font
    ws1.cell(row=1, column=1).fill = header_fill
    ws1.cell(row=1, column=2, value="Value").font = header_font
    ws1.cell(row=1, column=2).fill = header_fill

    param_items = [
        ("model_variant", getattr(args, "model_variant", None) or "v1"),
        ("rgb_backbone", getattr(args, "rgb_backbone", None) or "simple_cnn"),
        ("global_backbone", getattr(args, "global_backbone", None) or "simple_cnn"),
        ("race", getattr(args, "race", None) or "N/A"),
        ("no_pretrained", bool(getattr(args, "no_pretrained", False))),
        ("data_root", getattr(args, "data_root", None) or "N/A"),
        ("gt_excel", getattr(args, "gt_excel", None) or "N/A"),
        ("output_root", getattr(args, "output_root", None) or "N/A"),
    ]
    for r, (k, v) in enumerate(param_items, 2):
        ws1.cell(row=r, column=1, value=k).border = thin_border
        ws1.cell(row=r, column=2, value=str(v)).border = thin_border

    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 50

    # ==================== Sheet 2: Overall ====================
    ws2 = wb.create_sheet("Overall")
    write_header(ws2, ["Metric", "Value"])
    for r, (k, v) in enumerate(overall.items(), 2):
        ws2.cell(row=r, column=1, value=k).border = thin_border
        cell = ws2.cell(row=r, column=2, value=round(v, 6) if isinstance(v, float) else v)
        cell.border = thin_border
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 16

    # ==================== Sheet 3: PerGroup ====================
    ws3 = wb.create_sheet("PerGroup")
    group_headers = ["Model", "iOr", "Scene", "n_samples", "Pearson_r", "MAE"]
    write_header(ws3, group_headers)

    # 颜色条件格式：pearson_r >= 0.8 绿色，< 0 红色
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="F4B4C2", end_color="F4B4C2", fill_type="solid")

    for r, g in enumerate(per_group, 2):
        row_data = [g["model"], g["ior"], g["scene"], g["n_samples"], g["pearson_r"], g["mae"]]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws3.cell(row=r, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx <= 4 else "right")
        # pearson_r 列着色
        pr = g["pearson_r"]
        if not np.isnan(pr):
            pearson_cell = ws3.cell(row=r, column=5)
            if pr >= 0.8:
                pearson_cell.fill = green_fill
            elif pr < 0:
                pearson_cell.fill = red_fill

    col_widths = [10, 8, 12, 12, 14, 12]
    for i, w in enumerate(col_widths, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"

    # ==================== Sheet 4: RawData ====================
    ws4 = wb.create_sheet("RawData")
    write_header(ws4, ["original_name", "scene_person", "model", "ior", "scene", "pred", "target", "abs_error"])
    for r, (meta, pred, tgt) in enumerate(zip(metadata, predictions, targets), 2):
        model, ior, scene = _extract_model_ior_scene(meta["original_name"])
        row_data = [
            meta["original_name"],
            meta["scene_person"],
            model, ior, scene,
            round(float(pred), 6),
            round(float(tgt), 6),
            round(abs(float(pred) - float(tgt)), 6),
        ]
        for col_idx, val in enumerate(row_data, 1):
            ws4.cell(row=r, column=col_idx, value=val).border = thin_border
    raw_widths = [28, 16, 8, 6, 10, 12, 12, 12]
    for i, w in enumerate(raw_widths, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = "A2"

    # 保存
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Results exported to: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Test-only inference with per-group pearson_r export")
    parser.add_argument("--model-variant", choices=["v1", "v2", "v3"], default=None)
    parser.add_argument("--rgb-backbone", choices=["simple_cnn", "resnet50", "efficientnet_b0",
                                                    "mobilenet_v3_small", "mobilenet_v3_large",
                                                    "vit_b_16", "swin_t", "clip_vit_b32"], default=None)
    parser.add_argument("--global-backbone", choices=["simple_cnn", "mobilenet_v3_small", "mobilenet_v3_large", "efficientnet_b0"], default=None)
    parser.add_argument("--freeze-backbone", action="store_true", default=False)
    parser.add_argument("--no-pretrained", action="store_true")
    parser.add_argument("--race", choices=["CA", "AS", "SA", "AF"], default=None)
    parser.add_argument("--data-root", type=str, default=None)
    parser.add_argument("--gt-excel", type=str, default=None)
    parser.add_argument("--output-root", type=str, default=None)
    parser.add_argument("--batch-size", type=int, default=None)
    parser.add_argument("--num-workers", type=int, default=0)
    parser.add_argument("--resume", type=str, default=None,
                        help="Path to checkpoint file (e.g., /path/to/best_model.pth). "
                             "If not provided, uses CHECKPOINT_DIR/best_model.pth by default.")
    parser.add_argument("--ablation-face-only", action="store_true", default=False)
    parser.add_argument("--ablation-fusion-type", type=str,
                        choices=["gated", "concat", "se_gated", "cross_attn"], default=None)
    parser.add_argument("--exp-name", type=str, default=None)
    args = parser.parse_args()

    config = Config.get_config_dict()

    if args.data_root:
        dr = Path(args.data_root)
        config["FACE_RGB_ROOT"] = str(dr / "rendered_face")
        config["FACE_UV_ROOT"] = str(dr / "rendered_face_uv")
        config["GLOBAL_RGB_ROOT"] = str(dr / "rendered_2max")
        config["GT_EXCEL_PATH"] = str(dr / "gt" / "Preference_non_model_gt.xlsx")
    if args.gt_excel:
        config["GT_EXCEL_PATH"] = str(Path(args.gt_excel))
    if args.output_root:
        config["OUTPUT_ROOT"] = str(Path(args.output_root))
    if args.model_variant:
        config["MODEL_VARIANT"] = args.model_variant
    if args.rgb_backbone:
        config["MODEL"]["face_stream"]["rgb_backbone"] = args.rgb_backbone
        if args.rgb_backbone == "simple_cnn":
            config["MODEL"]["face_stream"]["rgb_pretrained"] = False
    if args.global_backbone:
        config["MODEL"]["global_stream"]["backbone"] = args.global_backbone
        if args.global_backbone == "simple_cnn":
            config["MODEL"]["global_stream"]["pretrained"] = False
    # Phase 0: freeze_backbone
    if bool(args.freeze_backbone):
        config["MODEL"]["face_stream"]["freeze_backbone"] = True
    else:
        _fn_backbone = str(config["MODEL"]["face_stream"].get("rgb_backbone", ""))
        if _fn_backbone in ("vit_b_16", "swin_t", "clip_vit_b32"):
            config["MODEL"]["face_stream"]["freeze_backbone"] = True
        else:
            config["MODEL"]["face_stream"]["freeze_backbone"] = False
    # Phase 0: ablation params
    if bool(args.ablation_face_only):
        config["ABLATION_FACE_ONLY"] = True
    if args.ablation_fusion_type is not None:
        config["ABLATION_FUSION_TYPE"] = str(args.ablation_fusion_type)
    if args.no_pretrained:
        config["MODEL"]["face_stream"]["rgb_pretrained"] = False
        config["MODEL"]["global_stream"]["pretrained"] = False
    if args.batch_size:
        config["TRAINING"]["batch_size"] = int(args.batch_size)

    if "DATALOADER" not in config or not isinstance(config["DATALOADER"], dict):
        config["DATALOADER"] = {}
    if args.num_workers is not None:
        config["DATALOADER"]["num_workers"] = int(args.num_workers)

    # Phase 0: exp-name 子目录
    if args.exp_name is not None:
        exp_name = str(args.exp_name).strip()
        if exp_name:
            config["OUTPUT_ROOT"] = str(Path(config["OUTPUT_ROOT"]) / exp_name)

    _apply_model_variant_overrides(config)
    if args.race:
        _apply_race_overrides(config, args.race)

    logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s", force=True)

    trainer = Trainer(config)

    # ---- 运行推理 ----
    logger.info("Running inference on test set...")
    result = run_test_with_metadata(trainer, ckpt_path=args.resume)

    overall = result["overall"]
    predictions = result["predictions"]
    targets = result["targets"]
    metadata = result["metadata"]

    # ---- 整体结果 ----
    print("\n" + "=" * 50)
    print("  OVERALL TEST RESULTS")
    print("=" * 50)
    for k, v in overall.items():
        print(f"  {k}: {v:.4f}" if isinstance(v, float) else f"  {k}: {v}")
    print("=" * 50)

    # ---- 按 model+ior+scene 分组 ----
    per_group = compute_per_group_pearson(
        np.array(predictions), np.array(targets), metadata
    )

    print(f"\nPer-group Pearson r ({len(per_group)} groups):")
    print(f"  {'Model':<6} {'iOr':<4} {'Scene':<8} {'N':>5} {'Pearson_r':>10} {'MAE':>10}")
    print(f"  {'-'*50}")
    for g in per_group:
        pr_str = f"{g['pearson_r']:.4f}" if not np.isnan(g['pearson_r']) else "N/A"
        mae_str = f"{g['mae']:.4f}" if not np.isnan(g['mae']) else "N/A"
        print(f"  {g['model']:<6} {g['ior']:<4} {g['scene']:<8} {g['n_samples']:>5} {pr_str:>10} {mae_str:>10}")

    # ---- 导出 xlsx ----
    xlsx_path = Path(config["RESULT_DIR"]) / "test_results.xlsx"
    export_to_xlsx(
        str(xlsx_path), args, overall, per_group,
        metadata, predictions, targets,
    )

    print(f"\nXLSX saved to: {xlsx_path}")


if __name__ == "__main__":
    main()
