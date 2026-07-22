"""
单独的测试/推理脚本 —— 加载 best_model.pth 在 test set 上跑评估，输出 xlsx。
支持 multi-head 模式：输出每个属性（10个）的 Overall / PerGroup / RawData sheet。

用法:
  python predict_p/test_only.py --model-variant v3 \
      --rgb-backbone simple_cnn --global-backbone simple_cnn --race SA \
      --multi-head --attributes all \
      --data-root /root/autodl-tmp \
      --gt-excel /root/autodl-tmp/gt/toMax_gt.xlsx \
      --output-root /root/autodl-tmp/deepskin/predict_p/output

输出 xlsx 包含:
  - Params sheet
  - Overall / Overall_{attr} sheet (每个属性)
  - PerGroup / PerGroup_{attr} sheet
  - RawData / RawData_{attr} sheet
"""
from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np
import pandas as pd
import torch
from torch.cuda.amp import autocast

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from predict_p.config import Config
from predict_p.trainer import Trainer
from predict_p.train import _apply_model_variant_overrides, _apply_race_overrides
from predict_p.metrics import evaluate_scores

logger = logging.getLogger(__name__)

# 10 属性名称（与 train.py 一致）
ALL_ATTRIBUTES = [
    "01Preference", "02Attractiveness", "03Feminine",
    "04Cooperative", "05Youth", "06Healthy",
    "07Fidelity", "08Harmony", "09Fair", "10Ruddy",
]


def _extract_model_ior_scene(original_name: str) -> tuple:
    s = str(original_name)
    parts = s.rsplit("_", 1)
    prefix = parts[0] if len(parts) == 2 and parts[1].isdigit() else s
    if len(prefix) >= 4:
        model = prefix[:3]
        ior = prefix[3]
        scene = prefix[4:] if len(prefix) > 4 else ""
    else:
        model = prefix; ior = ""; scene = ""
    return model, ior, scene


def _maybe_normalize_target_score(target_score: torch.Tensor) -> torch.Tensor:
    if target_score.max() > 1.5:
        return target_score / 10.0
    return target_score


def load_per_attribute_gt(gt_dir: str) -> tuple:
    """
    从 toMax_gt_{attr}.xlsx 加载每个属性的 GT 分数和 Lab 中心。
    返回: (name_to_scores, name_to_lab)
      - name_to_scores: {original_name: {attr_name: score}}
      - name_to_lab:    {original_name: {attr_name: (L, a, b)}}
    """
    gt_dir = Path(gt_dir)
    name_to_scores: Dict[str, Dict[str, float]] = {}
    name_to_lab: Dict[str, Dict[str, tuple]] = {}

    for attr in ALL_ATTRIBUTES:
        fpath = gt_dir / f"toMax_gt_{attr}.xlsx"
        if not fpath.exists():
            logger.warning(f"Per-attribute GT file not found: {fpath}")
            continue
        xls = pd.ExcelFile(fpath)
        total = 0
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name, header=0)
            if df.shape[1] < 5:
                continue
            for _, row in df.iterrows():
                name = str(row.iloc[0]).strip()
                try:
                    score = float(row.iloc[1])
                    L = float(row.iloc[2]); a = float(row.iloc[3]); b = float(row.iloc[4])
                except (ValueError, TypeError):
                    continue
                if name not in name_to_scores:
                    name_to_scores[name] = {}
                    name_to_lab[name] = {}
                name_to_scores[name][attr] = score
                # (50,0,0) = placeholder → NaN
                if abs(L-50)<0.01 and abs(a)<0.01 and abs(b)<0.01:
                    L=a=b=np.nan
                name_to_lab[name][attr] = (L, a, b)
                total += 1
        logger.info(f"  [{attr}] {total} scores from {len(xls.sheet_names)} sheets")

    return name_to_scores, name_to_lab


def run_test_with_metadata(trainer: Trainer, ckpt_path: str = None, split: str = "test") -> Dict:
    """
    跑推理。multi-head 时保留全部 10 列预测。
    """
    if ckpt_path is not None:
        ckpt_path = Path(ckpt_path)
    else:
        ckpt_path = Path(trainer.config["CHECKPOINT_DIR"]) / "best_model.pth"

    if ckpt_path.exists():
        ckpt = torch.load(ckpt_path, map_location=trainer.device)
        old_sd = ckpt["model_state_dict"]
        # Remap old checkpoint keys to current model naming:
        #   backbone -> backbone_raw  (face_stream.rgb_branch)
        #   score_head.heads -> score_heads
        new_sd = {}
        for k, v in old_sd.items():
            nk = k.replace("face_stream.rgb_branch.backbone.", "face_stream.rgb_branch.backbone_raw.")
            nk = nk.replace("score_head.heads.", "score_heads.")
            new_sd[nk] = v
        trainer.model.load_state_dict(new_sd, strict=True)
        logger.info(f"Loaded checkpoint from {ckpt_path}")
    else:
        logger.warning(f"Checkpoint not found: {ckpt_path}")

    trainer.model.eval()

    split = split.lower().strip()
    if split == "train":
        loader = trainer.train_loader
    elif split == "val":
        loader = trainer.val_loader
    else:
        loader = trainer.test_loader
    logger.info(f"Running inference on {split} set ({len(loader.dataset)} samples)")

    all_pred = []
    all_centers = []  # A9
    all_meta = []
    has_centers = False
    scaler_ctx = autocast() if trainer.scaler else torch.no_grad()

    with torch.no_grad():
        for batch in loader:
            face_rgb = batch["face_rgb"].to(trainer.device, non_blocking=True)
            global_rgb = batch.get("global_rgb", face_rgb).to(trainer.device, non_blocking=True)
            stat_features = batch.get("stat_features", None)
            if stat_features is not None:
                stat_features = stat_features.to(trainer.device, non_blocking=True)
            face_uv = None
            if str(trainer.config.get("MODEL_VARIANT", "v1")).lower().strip() != "v3":
                face_uv = batch["face_uv"].to(trainer.device, non_blocking=True)

            with scaler_ctx:
                model_out = trainer.model(face_rgb, face_uv, global_rgb, stat_features=stat_features)

            # A9: handle tuple (scores, centers)
            if isinstance(model_out, tuple):
                pred_logits, pred_centers = model_out
                has_centers = True
                all_centers.append(pred_centers.cpu())
            else:
                pred_logits = model_out

            all_pred.append(torch.sigmoid(pred_logits).cpu())

    # 从 loader.dataset 收集 metadata
    items = loader.dataset.data_items
    for item in items:
        all_meta.append({
            "original_name": str(item.get("original_name", "")),
            "scene_person": str(item.get("scene_person", "")),
        })

    ps = torch.cat(all_pred, dim=0).numpy()
    n = min(ps.shape[0], len(all_meta))
    ps = ps[:n]
    all_meta = all_meta[:n]

    result = {"predictions": ps, "metadata": all_meta}
    if has_centers:
        cs = torch.cat(all_centers, dim=0).numpy()[:n]  # (N, 3*10)
        result["centers"] = cs
    return result


def compute_metrics_for_column(pred_col: np.ndarray, tgt_col: np.ndarray) -> Dict[str, float]:
    """对单列 pred/tgt 计算指标。移除 NaN GT。"""
    mask = np.isfinite(tgt_col)
    if not mask.any():
        return {"mae": float("nan"), "rmse": float("nan"), "r2": float("nan"),
                "pearson_r": float("nan")}
    p = pred_col[mask]
    t = tgt_col[mask]
    return evaluate_scores(p, t)


def compute_deltae_ciede2000(lab_pred: np.ndarray, lab_gt: np.ndarray) -> np.ndarray:
    """
    Compute CIEDE2000 between predicted and GT Lab values.
    lab_pred: (N, 3) or list of (L,a,b) — predicted
    lab_gt:   (N, 3) or list of (L,a,b) — ground truth
    Returns: (N,) array of deltaE2000 values, NaN where GT is invalid.
    """
    try:
        from skimage.color import deltaE_ciede2000
    except ImportError:
        # Fallback: simple Euclidean deltaE
        dL = lab_pred[:, 0] - lab_gt[:, 0]
        da = lab_pred[:, 1] - lab_gt[:, 1]
        db = lab_pred[:, 2] - lab_gt[:, 2]
        return np.sqrt(dL**2 + da**2 + db**2)
    lab_pred = np.atleast_2d(lab_pred)
    lab_gt = np.atleast_2d(lab_gt)
    if lab_pred.shape != lab_gt.shape:
        raise ValueError(f"Shape mismatch: pred {lab_pred.shape}, gt {lab_gt.shape}")
    de = np.full(lab_pred.shape[0], np.nan)
    for i in range(lab_pred.shape[0]):
        if np.isfinite(lab_gt[i]).all():
            de[i] = deltaE_ciede2000(lab_pred[i], lab_gt[i])
    return de


def write_deltae_sheets(wb, suffix, pred_lab_arr, gt_lab_arr, metadata, header_font, header_fill, thin_border):
    """Add deltaE2000 Overall + PerGroup + RawData sheets."""
    import openpyxl as xl

    de = compute_deltae_ciede2000(pred_lab_arr, gt_lab_arr)
    valid = np.isfinite(de)
    if not valid.any():
        return

    def write_header(ws, headers, row=1):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=ci, value=h)
            c.font = header_font; c.fill = header_fill; c.border = thin_border

    # Overall deltaE
    ws_o = wb.create_sheet(f"DeltaE{suffix}")
    de_mean = np.mean(de[valid])
    de_med = np.median(de[valid])
    de_max = np.max(de[valid])
    de_min = np.min(de[valid])
    write_header(ws_o, ["Metric", "Value"])
    for ri, (k, v) in enumerate([("mean_DE2000", de_mean), ("median_DE2000", de_med),
                                  ("max_DE2000", de_max), ("min_DE2000", de_min),
                                  ("n_valid", int(valid.sum())), ("n_total", len(de))], 2):
        ws_o.cell(row=ri, column=1, value=k).border = thin_border
        ws_o.cell(row=ri, column=2, value=round(v, 6) if isinstance(v, float) else v).border = thin_border
    ws_o.column_dimensions["A"].width = 20
    ws_o.column_dimensions["B"].width = 16

    # PerGroup deltaE
    ws_p = wb.create_sheet(f"DeltaE_PerGroup{suffix}")
    write_header(ws_p, ["Model", "iOr", "Scene", "n_samples", "mean_DE2000"])
    groups = {}
    for i, meta in enumerate(metadata):
        if not np.isfinite(de[i]):
            continue
        model, ior, scene = _extract_model_ior_scene(meta["original_name"])
        key = f"{model}_{ior}_{scene}"
        if key not in groups:
            groups[key] = {"model": model, "ior": ior, "scene": scene, "de": []}
        groups[key]["de"].append(de[i])
    for ri, (key, g) in enumerate(sorted(groups.items()), 2):
        ws_p.cell(row=ri, column=1, value=g["model"]).border = thin_border
        ws_p.cell(row=ri, column=2, value=g["ior"]).border = thin_border
        ws_p.cell(row=ri, column=3, value=g["scene"]).border = thin_border
        ws_p.cell(row=ri, column=4, value=len(g["de"])).border = thin_border
        ws_p.cell(row=ri, column=5, value=round(np.mean(g["de"]), 6)).border = thin_border
    for i, w in enumerate([8, 6, 10, 10, 14], 1):
        ws_p.column_dimensions[xl.utils.get_column_letter(i)].width = w

    # RawData deltaE
    ws_r = wb.create_sheet(f"DeltaE_RawData{suffix}")
    write_header(ws_r, ["original_name", "scene_person", "model", "ior", "scene", "DE2000"])
    for ri, (meta, d) in enumerate(zip(metadata, de), 2):
        model, ior, scene = _extract_model_ior_scene(meta["original_name"])
        row_data = [meta["original_name"], meta["scene_person"], model, ior, scene,
                    round(float(d), 6) if np.isfinite(d) else "N/A"]
        for ci, val in enumerate(row_data, 1):
            ws_r.cell(row=ri, column=ci, value=val).border = thin_border
    for i, w in enumerate([28, 16, 8, 6, 10, 14], 1):
        ws_r.column_dimensions[xl.utils.get_column_letter(i)].width = w
    ws_r.freeze_panes = "A2"


def compute_per_group_pearson(ps: np.ndarray, ts: np.ndarray, metadata: list) -> List[Dict]:
    groups: Dict[str, Dict] = {}
    for i, meta in enumerate(metadata):
        if not np.isfinite(ts[i]):
            continue
        model, ior, scene = _extract_model_ior_scene(meta["original_name"])
        key = f"{model}_{ior}_{scene}"
        if key not in groups:
            groups[key] = {"model": model, "ior": ior, "scene": scene, "pred": [], "tgt": []}
        groups[key]["pred"].append(float(ps[i]))
        groups[key]["tgt"].append(float(ts[i]))

    results = []
    for key in sorted(groups.keys()):
        g = groups[key]
        pa, ta = np.array(g["pred"]), np.array(g["tgt"])
        n = len(pa)
        pearson_r = float(np.corrcoef(pa, ta)[0, 1]) if n > 1 else float("nan")
        mae = float(np.mean(np.abs(pa - ta))) if n > 0 else float("nan")
        results.append({"model": g["model"], "ior": g["ior"], "scene": g["scene"],
                         "n_samples": n, "pearson_r": pearson_r, "mae": mae})
    return results


def export_to_xlsx(
    output_path: str,
    args: argparse.Namespace,
    result: Dict,
    per_attr_gt: Optional[Dict[str, Dict[str, float]]] = None,
    per_attr_lab: Optional[Dict[str, Dict[str, tuple]]] = None,
) -> None:
    """导出 xlsx，multi-head 时每属性独立 sheet。"""
    import openpyxl as xl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = xl.Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"),
                          top=Side(style="thin"), bottom=Side(style="thin"))

    def write_header(ws, headers, row=1):
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = header_align; cell.border = thin_border

    def write_overall_sheet(ws, metrics):
        write_header(ws, ["Metric", "Value"])
        for r, (k, v) in enumerate(metrics.items(), 2):
            ws.cell(row=r, column=1, value=k).border = thin_border
            ws.cell(row=r, column=2,
                    value=round(v, 6) if isinstance(v, float) else v).border = thin_border
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 16

    def write_pergroup_sheet(ws, per_group):
        headers = ["Model", "iOr", "Scene", "n_samples", "Pearson_r", "MAE"]
        write_header(ws, headers)
        for r, g in enumerate(per_group, 2):
            ws.cell(row=r, column=1, value=g["model"]).border = thin_border
            ws.cell(row=r, column=2, value=g["ior"]).border = thin_border
            ws.cell(row=r, column=3, value=g["scene"]).border = thin_border
            ws.cell(row=r, column=4, value=g["n_samples"]).border = thin_border
            ws.cell(row=r, column=5,
                    value=round(g["pearson_r"], 6) if not np.isnan(g["pearson_r"]) else "N/A").border = thin_border
            ws.cell(row=r, column=6,
                    value=round(g["mae"], 6) if not np.isnan(g["mae"]) else "N/A").border = thin_border
        for i, w in enumerate([8, 6, 10, 10, 12, 12], 1):
            ws.column_dimensions[xl.utils.get_column_letter(i)].width = w

    def write_rawdata_sheet(ws, metadata, preds, tgts):
        headers = ["original_name", "scene_person", "model", "ior", "scene", "pred", "target", "abs_error"]
        write_header(ws, headers)
        for r, (meta, pred, tgt) in enumerate(zip(metadata, preds, tgts), 2):
            model, ior, scene = _extract_model_ior_scene(meta["original_name"])
            err = abs(float(pred) - float(tgt)) if np.isfinite(tgt) else float("nan")
            row_data = [meta["original_name"], meta["scene_person"], model, ior, scene,
                        round(float(pred), 6), round(float(tgt), 6),
                        round(err, 6) if np.isfinite(err) else "N/A"]
            for ci, val in enumerate(row_data, 1):
                ws.cell(row=r, column=ci, value=val).border = thin_border
        for i, w in enumerate([28, 16, 8, 6, 10, 12, 12, 12], 1):
            ws.column_dimensions[xl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ps = result["predictions"]  # (N, C) or (N,)
    metadata = result["metadata"]

    is_multi = ps.ndim == 2 and ps.shape[1] > 1
    attrs = ALL_ATTRIBUTES if is_multi else ["01Preference"]

    # ===== Params =====
    ws1 = wb.active; ws1.title = "Params"
    ws1.cell(row=1, column=1, value="Parameter").font = header_font
    ws1.cell(row=1, column=1).fill = header_fill
    ws1.cell(row=1, column=2, value="Value").font = header_font
    ws1.cell(row=1, column=2).fill = header_fill
    param_items = [
        ("model_variant", getattr(args, "model_variant", None) or "v1"),
        ("rgb_backbone", getattr(args, "rgb_backbone", None) or "simple_cnn"),
        ("global_backbone", getattr(args, "global_backbone", None) or "simple_cnn"),
        ("race", getattr(args, "race", None) or "N/A"),
        ("split", getattr(args, "split", "test")),
        ("multi_head", bool(getattr(args, "multi_head", False))),
        ("attributes", getattr(args, "attributes", None) or "N/A"),
        ("data_root", getattr(args, "data_root", None) or "N/A"),
        ("gt_excel", getattr(args, "gt_excel", None) or "N/A"),
        ("output_root", getattr(args, "output_root", None) or "N/A"),
    ]
    for r, (k, v) in enumerate(param_items, 2):
        ws1.cell(row=r, column=1, value=k).border = thin_border
        ws1.cell(row=r, column=2, value=str(v)).border = thin_border
    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 50

    # ===== Per-attribute sheets =====
    has_centers = "centers" in result and per_attr_lab is not None
    for col_idx, attr in enumerate(attrs):
        suffix = f"_{attr}"
        pred_col = ps[:, col_idx] if is_multi else ps.flatten()

        # Load GT score for this attribute
        tgt_col = np.full(len(metadata), np.nan)
        if per_attr_gt and is_multi:
            for i, meta in enumerate(metadata):
                name = meta["original_name"]
                tgt_col[i] = per_attr_gt.get(name, {}).get(attr, np.nan)

        # Compute metrics
        metrics = compute_metrics_for_column(pred_col, tgt_col)
        write_overall_sheet(wb.create_sheet(f"Overall{suffix}"), metrics)
        per_group = compute_per_group_pearson(pred_col, tgt_col, metadata)
        write_pergroup_sheet(wb.create_sheet(f"PerGroup{suffix}"), per_group)
        write_rawdata_sheet(wb.create_sheet(f"RawData{suffix}"), metadata,
                            pred_col.tolist(), tgt_col.tolist())

        # deltaE2000 sheets (A9/A10 only)
        if has_centers:
            centers = result["centers"]  # (N, 30)
            c_start = col_idx * 3
            pred_lab = centers[:, c_start:c_start+3]  # (N, 3)
            gt_lab = np.full((len(metadata), 3), np.nan)
            for i, meta in enumerate(metadata):
                name = meta["original_name"]
                lab = per_attr_lab.get(name, {}).get(attr, (np.nan, np.nan, np.nan))
                gt_lab[i] = lab
            write_deltae_sheets(wb, suffix, pred_lab, gt_lab, metadata,
                                header_font, header_fill, thin_border)

    # ===== Legacy aggregate sheets (first attr only, for backward compat) =====
    pred0 = ps[:, 0] if is_multi else ps.flatten()
    tgt0 = np.full(len(metadata), np.nan)
    if per_attr_gt and is_multi:
        for i, meta in enumerate(metadata):
            tgt0[i] = per_attr_gt.get(meta["original_name"], {}).get(attrs[0], np.nan)
    elif not is_multi:
        tgt0 = pred0  # placeholder, single-head GT not separately tracked here

    write_overall_sheet(wb.create_sheet("Overall"),
                        compute_metrics_for_column(pred0, tgt0))
    write_pergroup_sheet(wb.create_sheet("PerGroup"),
                         compute_per_group_pearson(pred0, tgt0, metadata))
    write_rawdata_sheet(wb.create_sheet("RawData"), metadata,
                        pred0.tolist(), tgt0.tolist())

    # Save
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Results exported to: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Test-only inference with per-attribute export")
    parser.add_argument("--model-variant", choices=["v1", "v2", "v3"], default=None)
    parser.add_argument("--rgb-backbone",
                        choices=["simple_cnn", "resnet50", "efficientnet_b0",
                                 "mobilenet_v3_small", "mobilenet_v3_large",
                                 "vit_b_16", "swin_t", "clip_vit_b32"], default=None)
    parser.add_argument("--global-backbone",
                        choices=["simple_cnn", "mobilenet_v3_small", "mobilenet_v3_large",
                                 "efficientnet_b0"], default=None)
    parser.add_argument("--freeze-backbone", action="store_true", default=False)
    parser.add_argument("--no-pretrained", action="store_true")
    parser.add_argument("--race", choices=["CA", "AS", "SA", "AF"], default=None)
    parser.add_argument("--split", choices=["train", "val", "test"], default="test")
    parser.add_argument("--data-root", type=str, default=None)
    parser.add_argument("--gt-excel", type=str, default=None)
    parser.add_argument("--output-root", type=str, default=None)
    parser.add_argument("--batch-size", type=int, default=None)
    parser.add_argument("--num-workers", type=int, default=0)
    parser.add_argument("--resume", type=str, default=None)
    parser.add_argument("--ablation-face-only", action="store_true", default=False)
    parser.add_argument("--ablation-fusion-type", type=str,
                        choices=["gated", "concat", "se_gated", "cross_attn"], default=None)
    parser.add_argument("--ablation-stat-stream", action="store_true", default=False)
    parser.add_argument("--ablation-lab-center", action="store_true", default=False)
    parser.add_argument("--ablation-loss-type", type=str, choices=["mse", "bce_smooth_l1"], default=None)
    parser.add_argument("--multi-head", action="store_true", default=False)
    parser.add_argument("--attributes", type=str, default=None)
    parser.add_argument("--exp-name", type=str, default=None)
    args = parser.parse_args()

    config = Config.get_config_dict()
    if args.data_root:
        dr = Path(args.data_root)
        config["FACE_RGB_ROOT"] = str(dr / "rendered_face")
        config["FACE_UV_ROOT"] = str(dr / "rendered_face_uv")
        config["GLOBAL_RGB_ROOT"] = str(dr / "rendered_2max")
        config["GT_EXCEL_PATH"] = str(dr / "gt" / "Preference_non_model_gt.xlsx")
    if args.gt_excel:
        config["GT_EXCEL_PATH"] = str(Path(args.gt_excel))
    if args.output_root:
        config["OUTPUT_ROOT"] = str(Path(args.output_root))
    if args.model_variant:
        config["MODEL_VARIANT"] = args.model_variant
    if args.rgb_backbone:
        config["MODEL"]["face_stream"]["rgb_backbone"] = args.rgb_backbone
        if args.rgb_backbone == "simple_cnn":
            config["MODEL"]["face_stream"]["rgb_pretrained"] = False
    if args.global_backbone:
        config["MODEL"]["global_stream"]["backbone"] = args.global_backbone
        if args.global_backbone == "simple_cnn":
            config["MODEL"]["global_stream"]["pretrained"] = False
    if bool(args.freeze_backbone):
        config["MODEL"]["face_stream"]["freeze_backbone"] = True
    else:
        fb = str(config["MODEL"]["face_stream"].get("rgb_backbone", ""))
        config["MODEL"]["face_stream"]["freeze_backbone"] = fb in ("vit_b_16", "swin_t", "clip_vit_b32")
    if bool(args.ablation_face_only):
        config["ABLATION_FACE_ONLY"] = True
    if args.ablation_fusion_type is not None:
        config["ABLATION_FUSION_TYPE"] = str(args.ablation_fusion_type)
    if bool(args.ablation_stat_stream):
        config["ABLATION_STAT_STREAM"] = True
    if bool(args.ablation_lab_center):
        config["ABLATION_LAB_CENTER"] = True
    if args.ablation_loss_type is not None:
        config["ABLATION_LOSS_TYPE"] = str(args.ablation_loss_type)
    if bool(args.multi_head):
        config["MULTI_HEAD"] = True
        if args.attributes is None or str(args.attributes).strip().lower() == "all":
            config["ATTRIBUTE_HEAD_NAMES"] = list(ALL_ATTRIBUTES)
        else:
            config["ATTRIBUTE_HEAD_NAMES"] = [s.strip() for s in str(args.attributes).split(",") if s.strip()]
    if args.no_pretrained:
        config["MODEL"]["face_stream"]["rgb_pretrained"] = False
        config["MODEL"]["global_stream"]["pretrained"] = False
    if args.batch_size:
        config["TRAINING"]["batch_size"] = int(args.batch_size)
    if "DATALOADER" not in config or not isinstance(config["DATALOADER"], dict):
        config["DATALOADER"] = {}
    if args.num_workers is not None:
        config["DATALOADER"]["num_workers"] = int(args.num_workers)
    if args.exp_name is not None:
        en = str(args.exp_name).strip()
        if en:
            config["OUTPUT_ROOT"] = str(Path(config["OUTPUT_ROOT"]) / en)

    _apply_model_variant_overrides(config)
    if args.race:
        _apply_race_overrides(config, args.race)

    logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s", force=True)

    trainer = Trainer(config)

    # 加载 per-attribute GT（multi-head 时）
    per_attr_gt = None
    per_attr_lab = None
    if bool(args.multi_head) and args.data_root:
        gt_dir = Path(args.data_root) / "gt"
        logger.info(f"Loading per-attribute GT from {gt_dir} ...")
        per_attr_gt, per_attr_lab = load_per_attribute_gt(str(gt_dir))
        logger.info(f"  Loaded GT for {len(ALL_ATTRIBUTES)} attributes × {len(per_attr_gt)} samples")

    # 推理
    logger.info(f"Running inference on {args.split} set...")
    result = run_test_with_metadata(trainer, ckpt_path=args.resume, split=args.split)

    # 打印 Overall (01Preference)
    ps = result["predictions"]
    pred0 = ps[:, 0] if (ps.ndim == 2 and ps.shape[1] > 1) else ps.flatten()
    metadata = result["metadata"]
    tgt0 = np.full(len(metadata), np.nan)
    if per_attr_gt:
        for i, meta in enumerate(metadata):
            tgt0[i] = per_attr_gt.get(meta["original_name"], {}).get("01Preference", np.nan)
    m0 = compute_metrics_for_column(pred0, tgt0)
    print("\n" + "=" * 50)
    print("  OVERALL (01Preference)")
    print("=" * 50)
    for k, v in m0.items():
        print(f"  {k}: {v:.4f}" if isinstance(v, float) else f"  {k}: {v}")

    if bool(args.multi_head):
        print("\n  Per-attribute Pearson r:")
        for col_idx, attr in enumerate(ALL_ATTRIBUTES):
            pc = ps[:, col_idx]
            tc = np.full(len(metadata), np.nan)
            if per_attr_gt:
                for i, meta in enumerate(metadata):
                    tc[i] = per_attr_gt.get(meta["original_name"], {}).get(attr, np.nan)
            am = compute_metrics_for_column(pc, tc)
            r_str = f"{am.get('pearson_r', float('nan')):.4f}" if np.isfinite(am.get('pearson_r', float('nan'))) else "N/A"
            print(f"    {attr}: pearson_r={r_str}")

    print("=" * 50)

    # 导出 xlsx
    exp_dir = Path(trainer.config["OUTPUT_DIR"])
    output_xlsx = exp_dir / "results" / f"test_results.xlsx"
    export_to_xlsx(str(output_xlsx), args, result, per_attr_gt, per_attr_lab)


if __name__ == "__main__":
    main()
