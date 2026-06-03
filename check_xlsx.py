"""检查推理结果 xlsx"""
import openpyxl, sys

path = sys.argv[1]
wb = openpyxl.load_workbook(path)

print("=== Overall ===")
ws = wb["Overall"]
for r in range(1, ws.max_row + 1):
    print(f"  {ws.cell(r, 1).value}: {ws.cell(r, 2).value}")

print("\n=== PerGroup (前20行) ===")
ws = wb["PerGroup"]
for r in range(1, min(ws.max_row + 1, 21)):
    vals = [ws.cell(r, c).value for c in range(1, 7)]
    print(f"  {vals}")

print("\n=== RawData (前10行) ===")
ws = wb["RawData"]
for r in range(1, min(ws.max_row + 1, 11)):
    vals = [ws.cell(r, c).value for c in range(1, 9)]
    print(f"  {vals}")
