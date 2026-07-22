"""
可复用数据完整性检查脚本 — check_data_integrity.py
==================================================
复用 predict_p.data.dataset 中的 GT 解析逻辑, 与训练保持一致。

功能:
  ③ GT 完整性矩阵: 每个 {model}{i/r} 的每个 scene, 对每个属性检查 preference_score 是否齐全
  ④ 修复 Lab=(50,0,0) → NaN (preference_score 为空时)
  ⑤ 四路交叉引用: rendered_2max ↔ rendered_face ↔ mask ↔ rendered_face_uv

输出:
  check_data_integrity_{serial}.xlsx  (单文件, 包含 Summary + 40 个 subfolder sheet)

用法 (云端):
  cd /root/autodl-tmp/deepskin
  python3 scripts/check_data_integrity.py --serial v4 --data-root /root/autodl-tmp --output-dir /root/autodl-tmp
"""

import argparse
import os
import re
import shutil
import sys
from pathlib import Path
from collections import defaultdict
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── 尝试复用 predict_p 中的 Dataset 逻辑 ───
try:
    _repo_root = Path(__file__).resolve().parents[1]
    if str(_repo_root) not in sys.path:
        sys.path.insert(0, str(_repo_root))
    from predict_p.data.dataset import FacialPreferenceDataset as _DS

    # 复用 Dataset 中的函数
    _parse_scene = staticmethod(_DS._extract_prefix_from_original_name)
    _extract_id = staticmethod(_DS._extract_id_from_name)
    # Lab=(50,0,0) NaN 检测逻辑完全照搬 _load_data 第166-167行
    _IS_LAB_PLACEHOLDER = lambda L, a, b: abs(L - 50.0) < 0.01 and abs(a) < 0.01 and abs(b) < 0.01
    _REUSE_DATASET = True
except Exception as e:
    print(f"[WARN] Cannot import predict_p.data.dataset: {e}")
    print("       Falling back to standalone parsing (consistent with training logic)")

    def _parse_scene(original_name: str) -> Optional[str]:
        """从 original_name 提取场景前缀, 与 Dataset._extract_prefix_from_original_name 一致"""
        s = str(original_name)
        parts = s.rsplit("_", 1)
        if len(parts) == 2 and parts[1].isdigit():
            return parts[0]
        return s

    def _extract_id(name: str) -> Optional[str]:
        m = re.match(r'^[fm](\d{2})[ir]$', str(name))
        return m.group(1) if m else None

    _IS_LAB_PLACEHOLDER = lambda L, a, b: abs(L - 50.0) < 0.01 and abs(a) < 0.01 and abs(b) < 0.01
    _REUSE_DATASET = False


# ═════════════════════════════════════════════════════════════════════===
# Constants
# ═════════════════════════════════════════════════════════════════════===
ATTR_SERIALS = [
    "01Preference", "02Attractiveness", "03Feminine", "04Cooperative",
    "05Youth", "06Healthy", "07Fidelity", "08Harmony", "09Fair", "10Ruddy",
]

# 40 subfolders
_ALL_MODELS = [f"{g}{n:02d}{l}" for g in ["f", "m"] for n in range(1, 11) for l in ["i", "r"]]

# Scene counts per location type
_I_SCENES = ["h3k", "h4k", "h5k", "h6k", "h7k", "h8k", "hd65",
             "l3k", "l4k", "l5k", "l6k", "l7k", "l8k", "ld65",
             "m3k", "m4k", "m5k", "m6k", "m7k", "m8k", "md65"]
_R_SCENES = [f"rs{s:02d}" for s in range(1, 15)]
EXPECTED_I_COUNT = len(_I_SCENES) * 33   # 21 × 33 = 693
EXPECTED_R_COUNT = len(_R_SCENES) * 33   # 14 × 33 = 462

# Styles
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
NA_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _safe_cell(row, idx):
    if idx < len(row):
        return row[idx]
    return None


# ═════════════════════════════════════════════════════════════════════===
# Task ④: Fix Lab=(50,0,0) for empty preference_score
# ═════════════════════════════════════════════════════════════════════===
def fix_lab_values(gt_dir: Path, dry_run: bool = False) -> int:
    """修复: preference_score 为空且 Lab=(50,0,0) → Lab 改为 NaN。
    与 Dataset._load_data 第166-167行的逻辑完全一致。
    """
    fixed_total = 0
    backup_done = False

    for attr in ATTR_SERIALS:
        xlsx_path = gt_dir / f"toMax_gt_{attr}.xlsx"
        if not xlsx_path.exists():
            print(f"  [SKIP] {xlsx_path.name} not found")
            continue

        if not dry_run and not backup_done:
            backup_dir = gt_dir / "backup"
            backup_dir.mkdir(exist_ok=True)
            for a in ATTR_SERIALS:
                src = gt_dir / f"toMax_gt_{a}.xlsx"
                if src.exists():
                    shutil.copy2(src, backup_dir / f"toMax_gt_{a}.xlsx")
            print(f"  Backed up to {backup_dir}/")
            backup_done = True

        wb = openpyxl.load_workbook(xlsx_path)
        attr_fixed = 0

        for sn in wb.sheetnames:
            ws = wb[sn]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            try:
                pref_idx = headers.index("preference_score")
                L_idx = headers.index("L*")
                a_idx = headers.index("a*")
                b_idx = headers.index("b*")
            except ValueError:
                continue

            for row in ws.iter_rows(min_row=2):
                pref_cell = _safe_cell(row, pref_idx)
                L_cell = _safe_cell(row, L_idx)
                a_cell = _safe_cell(row, a_idx)
                b_cell = _safe_cell(row, b_idx)

                if pref_cell is None or L_cell is None:
                    continue

                pref_val = pref_cell.value
                L_val = L_cell.value
                a_val = a_cell.value if a_cell else None
                b_val = b_cell.value if b_cell else None

                # 检查 pref 为空
                pref_empty = (pref_val is None or
                              (isinstance(pref_val, str) and pref_val.strip() == "") or
                              (isinstance(pref_val, float) and (pref_val != pref_val)))  # NaN

                if pref_empty and all(v is not None for v in [L_val, a_val, b_val]):
                    try:
                        if _IS_LAB_PLACEHOLDER(float(L_val), float(a_val), float(b_val)):
                            if not dry_run:
                                L_cell.value = None
                                a_cell.value = None
                                b_cell.value = None
                            attr_fixed += 1
                            if attr_fixed <= 3:
                                print(f"  [{attr}/{sn}] Row {row[0].row}: Lab=(50,0,0)→NaN")
                    except (TypeError, ValueError):
                        continue

        if attr_fixed > 0:
            print(f"  [{attr}] Fixed {attr_fixed} rows")
            fixed_total += attr_fixed
            if not dry_run:
                wb.save(xlsx_path)
        wb.close()

    return fixed_total


# ═════════════════════════════════════════════════════════════════════===
# Task ③: GT completeness matrix — 复用 Dataset 的 GT 解析逻辑
# ═════════════════════════════════════════════════════════════════════===
def build_gt_completeness(gt_dir: Path) -> Dict[str, Dict[str, Dict[str, bool]]]:
    """
    返回: { model_ior: { attr_serial: { scene_prefix: True/False } } }

    逻辑与 Dataset._load_attribute_scores 相同:
    - 打开 toMax_gt_{attr}.xlsx
    - sheet = model_ior (e.g. "f01i")
    - col0 = original_name, col1 = preference_score
    - 通过 _parse_scene 提取 scene 前缀,
      检查每个 scene 的 33 行 preference_score 是否全部有值
    """
    completeness = defaultdict(lambda: defaultdict(dict))

    for attr in ATTR_SERIALS:
        xlsx_path = gt_dir / f"toMax_gt_{attr}.xlsx"
        if not xlsx_path.exists():
            print(f"  [SKIP] {xlsx_path.name} not found")
            continue

        wb = openpyxl.load_workbook(xlsx_path, read_only=True)

        for sn in wb.sheetnames:
            ws = wb[sn]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            try:
                pref_idx = headers.index("preference_score")
                name_idx = headers.index("original_name")
            except ValueError:
                continue

            # 按 scene 分组
            scene_rows = defaultdict(list)
            for row in ws.iter_rows(min_row=2, values_only=True):
                oname = _safe_cell(row, name_idx)
                pref = _safe_cell(row, pref_idx)
                scene = _parse_scene(oname)
                if scene:
                    scene_rows[scene].append(pref)

            for scene, prefs in scene_rows.items():
                # 与 Dataset 相同: preference_score 为 None/空字符串/NaN → 缺失
                all_ok = all(
                    p is not None and
                    not (isinstance(p, str) and p.strip() == "") and
                    not (isinstance(p, float) and (p != p))  # NaN
                    for p in prefs
                )
                completeness[sn][attr][scene] = all_ok

        wb.close()

    return dict(completeness)


# ═════════════════════════════════════════════════════════════════════===
# Task ⑤: Cross-reference — 与 Dataset 中构造路径的逻辑一致
# ═════════════════════════════════════════════════════════════════════===
def check_cross_reference(
    base_2max: Path, base_face: Path, base_mask: Path, base_uv: Path,
) -> Dict[str, Dict]:
    """
    检查每张 rendered_2max 图片是否有对应的:
      - rendered_face/{model}/{scene}_{idx:02d}_face.jpg
      - mask/{model}/{scene_short}.JPG (大小写不敏感)
      - rendered_face_uv/{model}/{scene}_{idx:02d}_face_uv.npy

    返回: { model_ior: { "count_ok": bool, "groups": { scene: {...} }, "overall_ok": bool } }
    """
    from PIL import Image

    results = {}

    for sf in sorted(_ALL_MODELS):
        sf_dir = base_2max / sf
        if not sf_dir.is_dir():
            print(f"  [SKIP] {sf_dir} not found")
            continue

        jpg_files = sorted(sf_dir.glob("*.jpg"))
        actual_count = len(jpg_files)
        loc_type = "I" if sf[-1] == "i" else "R"
        expected_count = EXPECTED_I_COUNT if loc_type == "I" else EXPECTED_R_COUNT
        count_ok = (actual_count == expected_count)

        # 按 scene 分组
        scene_groups: Dict[str, List[Path]] = defaultdict(list)
        for f in jpg_files:
            scene = _parse_scene(f.stem)
            if scene:
                scene_groups[scene].append(f)

        groups_result = {}
        all_groups_ok = True

        for scene, files in sorted(scene_groups.items()):
            n_in_group = len(files)
            all_33 = (n_in_group == 33)

            # 检查图片尺寸
            dims = set()
            for f in files:
                try:
                    img = Image.open(f)
                    dims.add(img.size)
                    img.close()
                except Exception:
                    pass
            dim_equal = (len(dims) == 1)

            # 四路交叉引用
            face_ok = True
            mask_ok = True
            uv_ok = True

            for idx in range(1, 34):
                stem = f"{scene}_{idx:02d}"

                # Face: {model}/{scene}_{idx:02d}_face.jpg
                if not (base_face / sf / f"{stem}_face.jpg").exists():
                    face_ok = False

                # Mask: {model}/{scene_short}.JPG (strip model prefix)
                scene_short = scene[len(sf):]
                mask_dir = base_mask / sf
                mask_found = any(
                    f.stem.lower() == scene_short.lower()
                    for f in (mask_dir.iterdir() if mask_dir.is_dir() else [])
                    if f.is_file()
                )
                if not mask_found:
                    mask_ok = False

                # UV: {model}/{scene}_{idx:02d}_face_uv.npy
                if not (base_uv / sf / f"{stem}_face_uv.npy").exists():
                    uv_ok = False

            group_ok = all_33 and dim_equal and face_ok and mask_ok and uv_ok
            if not group_ok:
                all_groups_ok = False

            groups_result[scene] = {
                "count": n_in_group,
                "dim_equal": dim_equal,
                "rendered2max": all_33,
                "rendered_face": face_ok,
                "mask": mask_ok,
                "uv": uv_ok,
                "all_ok": group_ok,
            }

        overall_ok = count_ok and all_groups_ok

        results[sf] = {
            "type": loc_type,
            "actual_count": actual_count,
            "expected_count": expected_count,
            "count_ok": count_ok,
            "groups": groups_result,
            "overall_ok": overall_ok,
        }
        status = "PASS" if overall_ok else "FAIL"
        print(f"  [{sf}] {actual_count}/{expected_count} imgs, {len(groups_result)} groups → {status}")

    return results


# ═════════════════════════════════════════════════════════════════════===
# Excel writer — 综合输出(③ GT 完整性 + ⑤ 交叉引用)
# ═════════════════════════════════════════════════════════════════════===
def _style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN_BORDER


def _auto_width(ws, min_w=10, max_w=28):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(min_w, min(max_len + 2, max_w))


def write_comprehensive_xlsx(
    completeness: dict,
    crossref: dict,
    output_path: Path,
    serial: str,
    fixed_count: int,
) -> None:
    """写综合 xlsx: Summary + 40 个 subfolder sheet (GT×CrossRef 合并)"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── Summary sheet ──
    ws_sum = wb.create_sheet("Summary")
    headers = [
        "Subfolder", "Type", "Rendered2max", "Face_All", "Mask_All",
        "UV_All", "GT_All_OK", "GT_Partial", "Count_OK", "Count",
        "Groups", "Groups_OK", "CrossRef_OK", "Overall",
    ]
    for c, h in enumerate(headers, 1):
        ws_sum.cell(row=1, column=c, value=h)
    _style_header(ws_sum, 1, len(headers))

    total_pass = 0
    row = 2
    for sf in sorted(_ALL_MODELS):
        gtc = completeness.get(sf, {})
        xr = crossref.get(sf, {})

        # GT 统计
        gt_scenes = set()
        for attr in ATTR_SERIALS:
            gt_scenes.update(gtc.get(attr, {}).keys())
        gt_total = len(gt_scenes) * len(ATTR_SERIALS)
        gt_ok = sum(
            1 for attr in ATTR_SERIALS
            for s, v in gtc.get(attr, {}).items() if v
        )
        gt_all_ok = (gt_ok == gt_total) if gt_total > 0 else None

        # Cross-ref 统计
        if xr:
            groups = xr.get("groups", {})
            face_all = all(g.get("rendered_face", False) for g in groups.values()) if groups else None
            mask_all = all(g.get("mask", False) for g in groups.values()) if groups else None
            uv_all = all(g.get("uv", False) for g in groups.values()) if groups else None
            r2m_all = all(g.get("rendered2max", False) for g in groups.values()) if groups else None
            groups_ok = sum(1 for g in groups.values() if g.get("all_ok")) if groups else 0
            xref_ok = xr.get("overall_ok", False)
            count_str = f"{xr.get('actual_count', '?')}/{xr.get('expected_count', '?')}"
            count_ok = xr.get("count_ok", False)
            loc_type = xr.get("type", "?")
            n_groups = len(groups)
        else:
            face_all = mask_all = uv_all = r2m_all = None
            groups_ok = 0
            xref_ok = False
            count_str = "N/A"
            count_ok = False
            loc_type = "?"
            n_groups = 0

        overall = (bool(gt_all_ok) and bool(xref_ok))

        vals = [
            sf, loc_type,
            "TRUE" if r2m_all else ("FALSE" if r2m_all is not None else "N/A"),
            "TRUE" if face_all else ("FALSE" if face_all is not None else "N/A"),
            "TRUE" if mask_all else ("FALSE" if mask_all is not None else "N/A"),
            "TRUE" if uv_all else ("FALSE" if uv_all is not None else "N/A"),
            "TRUE" if gt_all_ok else ("FALSE" if gt_all_ok is not None else "N/A"),
            f"{gt_ok}/{gt_total}" if gt_total > 0 else "N/A",
            "TRUE" if count_ok else "FALSE",
            count_str,
            n_groups,
            f"{groups_ok}/{n_groups}" if n_groups else "N/A",
            "PASS" if xref_ok else "FAIL",
            "PASS" if overall else "FAIL",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws_sum.cell(row=row, column=c, value=v)
            if isinstance(v, str):
                if v == "TRUE" or v == "PASS":
                    cell.fill = PASS_FILL
                elif v == "FALSE" or v == "FAIL":
                    cell.fill = FAIL_FILL
                elif v == "N/A":
                    cell.fill = NA_FILL
        if overall:
            total_pass += 1
        row += 1

    # Summary footer
    ws_sum.cell(row=row, column=1, value=f"Total PASS: {total_pass}/{len(_ALL_MODELS)}").font = Font(bold=True)
    ws_sum.cell(row=row + 1, column=1,
                value=f"Task ④ fixed: {fixed_count} rows (Lab=(50,0,0)→NaN)").font = Font(italic=True)
    ws_sum.cell(row=row + 2, column=1,
                value=f"Dataset logic reused: {'Yes' if _REUSE_DATASET else 'No (standalone fallback)'}").font = Font(italic=True)
    _auto_width(ws_sum, min_w=12)
    ws_sum.freeze_panes = "A2"

    # ── Per-subfolder sheets ──
    for sf in sorted(_ALL_MODELS):
        ws = wb.create_sheet(sf)

        gtc = completeness.get(sf, {})
        xr = crossref.get(sf, {})
        xr_groups = xr.get("groups", {}) if xr else {}

        # Collect all scenes across GT attrs + crossref
        all_scenes = set()
        for attr in ATTR_SERIALS:
            all_scenes.update(gtc.get(attr, {}).keys())
        all_scenes.update(xr_groups.keys())
        all_scenes = sorted(all_scenes)

        # Headers: Scene | GT_{attr}×10 | GT_All_OK | Rendered2max | RenderedFace | Mask | UV | CrossRef_OK | Overall
        headers = ["Scene"]
        for attr in ATTR_SERIALS:
            headers.append(attr[:2])  # abbreviated: "01", "02", ...
        headers += ["GT_All", "R2max", "Face", "Mask", "UV", "XRef_OK", "Overall"]
        for c, h in enumerate(headers, 1):
            ws.cell(row=1, column=c, value=h)
        _style_header(ws, 1, len(headers))

        for ri, scene in enumerate(all_scenes, 2):
            ws.cell(row=ri, column=1, value=scene)

            # GT columns
            gt_all_ok = True
            for ci, attr in enumerate(ATTR_SERIALS, 2):
                val = gtc.get(attr, {}).get(scene, None)
                if val is None:
                    c = ws.cell(row=ri, column=ci, value="N/A")
                    c.fill = NA_FILL
                    gt_all_ok = False
                elif val:
                    c = ws.cell(row=ri, column=ci, value="TRUE")
                    c.fill = PASS_FILL
                else:
                    c = ws.cell(row=ri, column=ci, value="FALSE")
                    c.fill = FAIL_FILL
                    gt_all_ok = False

            # GT All OK
            gt_col = 2 + len(ATTR_SERIALS)
            c = ws.cell(row=ri, column=gt_col, value="TRUE" if gt_all_ok else "FALSE")
            c.fill = PASS_FILL if gt_all_ok else FAIL_FILL

            # Cross-ref columns
            g = xr_groups.get(scene, {})
            for offset, key, label in [(1, "rendered2max", "TRUE"), (2, "rendered_face", "TRUE"),
                                        (3, "mask", "TRUE"), (4, "uv", "TRUE")]:
                v = g.get(key, False)
                c = ws.cell(row=ri, column=gt_col + offset,
                            value="TRUE" if v else "FALSE")
                c.fill = PASS_FILL if v else FAIL_FILL

            xref_scene_ok = g.get("all_ok", False)
            c = ws.cell(row=ri, column=gt_col + 5, value="TRUE" if xref_scene_ok else "FALSE")
            c.fill = PASS_FILL if xref_scene_ok else FAIL_FILL

            # Overall
            overall_scene = gt_all_ok and xref_scene_ok
            c = ws.cell(row=ri, column=gt_col + 6, value="PASS" if overall_scene else "FAIL")
            c.fill = PASS_FILL if overall_scene else FAIL_FILL
            c.font = Font(bold=True)

        _auto_width(ws, min_w=8, max_w=16)
        ws.freeze_panes = "B2"

    wb.save(output_path)
    print(f"\n✅ Comprehensive report saved: {output_path}")
    print(f"   Sheets: Summary + {len(_ALL_MODELS)} subfolders")


# ═════════════════════════════════════════════════════════════════════===
# Main
# ═════════════════════════════════════════════════════════════════════===
def main():
    parser = argparse.ArgumentParser(
        description="可复用数据完整性检查(复用 predict_p Dataset 逻辑)")
    parser.add_argument("--serial", default="v4", help="版本后缀")
    parser.add_argument("--data-root", default="/root/autodl-tmp",
                        help="数据根目录(含 rendered_2max/rendered_face/mask/rendered_face_uv/gt)")
    parser.add_argument("--output-dir", default=None,
                        help="输出目录(默认同 --data-root)")
    parser.add_argument("--skip-fix", action="store_true", help="跳过 Task ④")
    parser.add_argument("--dry-run", action="store_true", help="Task ④ dry-run")
    args = parser.parse_args()

    data_root = Path(args.data_root)
    output_dir = Path(args.output_dir) if args.output_dir else data_root

    gt_dir = data_root / "gt"
    base_2max = data_root / "rendered_2max"
    base_face = data_root / "rendered_face"
    base_mask = data_root / "mask"
    base_uv = data_root / "rendered_face_uv"

    print("=" * 60)
    print(f"Data integrity check — serial={args.serial}")
    print(f"Data root:     {data_root}")
    print(f"Output dir:    {output_dir}")
    print(f"Dataset reuse: {'✅ YES' if _REUSE_DATASET else '⚠️ standalone'}")
    print("=" * 60)

    # ── Task ④ ──
    fixed_count = 0
    if not args.skip_fix:
        print("\n📌 Task ④: Fix Lab=(50,0,0)→NaN for empty preference_score")
        print("-" * 40)
        fixed_count = fix_lab_values(gt_dir, dry_run=args.dry_run)
        print(f"  Total fixed: {fixed_count} rows")
    else:
        print("\n⏭️  Task ④: SKIPPED")

    # ── Task ③ ──
    print("\n📌 Task ③: GT completeness matrix")
    print("-" * 40)
    completeness = build_gt_completeness(gt_dir)
    total_cells = 0; ok_cells = 0
    for sf in _ALL_MODELS:
        for attr in ATTR_SERIALS:
            for scene, ok in completeness.get(sf, {}).get(attr, {}).items():
                total_cells += 1
                if ok: ok_cells += 1
    print(f"  GT completeness: {ok_cells}/{total_cells} ({100*ok_cells/max(total_cells,1):.1f}%)")

    # ── Task ⑤ ──
    print("\n📌 Task ⑤: Cross-reference check")
    print("-" * 40)
    crossref = check_cross_reference(base_2max, base_face, base_mask, base_uv)

    # ── Write comprehensive xlsx ──
    print("\n📌 Writing comprehensive report...")
    output_path = output_dir / f"check_data_integrity_{args.serial}.xlsx"
    write_comprehensive_xlsx(completeness, crossref, output_path, args.serial, fixed_count)

    # ── Quick summary ──
    xr_pass = sum(1 for v in crossref.values() if v.get("overall_ok"))
    gt_all_models_ok = 0
    for sf in _ALL_MODELS:
        gtc = completeness.get(sf, {})
        gt_scenes = set()
        for attr in ATTR_SERIALS:
            gt_scenes.update(gtc.get(attr, {}).keys())
        gt_total_sf = len(gt_scenes) * 10
        gt_ok_sf = sum(1 for attr in ATTR_SERIALS for s, v in gtc.get(attr, {}).items() if v)
        if gt_total_sf > 0 and gt_ok_sf == gt_total_sf:
            gt_all_models_ok += 1

    print("\n" + "=" * 60)
    print("📊 SUMMARY")
    print(f"  GT completeness: {gt_all_models_ok}/{len(_ALL_MODELS)} subfolders all-OK")
    print(f"  Cross-reference: {xr_pass}/{len(crossref)} subfolders PASS")
    print(f"  Lab fix (task ④): {fixed_count} rows changed")
    print(f"  Output: {output_path}")
    print("=" * 60)


if __name__ == "__main__":
    main()
